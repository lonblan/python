# ============================================================
# PASO 7: Leer alumnos desde un Excel y generar certificados
# ============================================================
# Tú creas (o editas) el archivo:  alumnos.xlsx
# Con dos columnas en la PRIMERA fila (encabezados):
#
#     | nombre        | curso   |
#     | Ana Pérez     | Python  |
#     | Luis Gómez    | Python  |
#
# El programa lee esa hoja y crea un PDF por cada fila.
#
# Librería nueva: openpyxl  (ya está en el .venv)
# ============================================================

from pathlib import Path
from fpdf import FPDF
from openpyxl import load_workbook  # Lee archivos .xlsx


class GeneradorCertificado:
    """Fábrica de certificados PDF a partir de un Excel."""

    def __init__(self):
        self.carpeta = Path(__file__).resolve().parent
        self.carpeta_salida = self.carpeta / "certificados_excel"
        self.carpeta_salida.mkdir(exist_ok=True)

        self.fuente_normal = "/System/Library/Fonts/Supplemental/Arial.ttf"
        self.fuente_negrita = "/System/Library/Fonts/Supplemental/Arial Bold.ttf"
        self.fuente_cursiva = "/System/Library/Fonts/Supplemental/Arial Italic.ttf"

        self.nombre = ""
        self.curso = ""
        self.fecha = ""

    def leer_excel(self, nombre_archivo="alumnos.xlsx"):
        """
        Abre el Excel y devuelve una lista de diccionarios:
        [{"nombre": "...", "curso": "..."}, ...]
        """
        ruta = self.carpeta / nombre_archivo

        if not ruta.exists():
            raise FileNotFoundError(
                f"No encuentro el archivo:\n{ruta}\n"
                "Pon alumnos.xlsx en la misma carpeta que este script."
            )

        libro = load_workbook(ruta)       # Abre el archivo
        hoja = libro.active               # Primera hoja

        # Fila 1 = encabezados. Buscamos las columnas "nombre" y "curso"
        encabezados = [celda.value for celda in hoja[1]]
        encabezados_norm = [
            (h or "").strip().lower() for h in encabezados
        ]

        try:
            col_nombre = encabezados_norm.index("nombre") + 1  # openpyxl es 1-based
            col_curso = encabezados_norm.index("curso") + 1
        except ValueError:
            raise ValueError(
                'La primera fila debe tener las columnas: nombre | curso'
            )

        alumnos = []
        # Desde la fila 2 hasta el final
        for fila in range(2, hoja.max_row + 1):
            nombre = hoja.cell(row=fila, column=col_nombre).value
            curso = hoja.cell(row=fila, column=col_curso).value

            # Saltar filas vacías
            if not nombre and not curso:
                continue

            alumnos.append({
                "nombre": str(nombre).strip(),
                "curso": str(curso).strip() if curso else "",
            })

        return alumnos

    def crear_pdf(self):
        """Arma el PDF en memoria y lo devuelve."""
        pdf = FPDF(orientation="L", unit="mm", format="A4")
        pdf.add_page()

        pdf.add_font("Arial", style="", fname=self.fuente_normal)
        pdf.add_font("Arial", style="B", fname=self.fuente_negrita)
        pdf.add_font("Arial", style="I", fname=self.fuente_cursiva)

        pdf.set_font("Arial", style="B", size=28)
        pdf.cell(w=0, h=20, text="CERTIFICADO DE CURSO", align="C")
        pdf.ln(30)

        pdf.set_font("Arial", style="", size=16)
        pdf.cell(w=0, h=10, text="Se otorga el presente certificado a:", align="C")
        pdf.ln(20)

        pdf.set_font("Arial", style="B", size=24)
        pdf.cell(w=0, h=12, text=self.nombre, align="C")
        pdf.ln(20)

        pdf.set_font("Arial", style="", size=16)
        pdf.cell(w=0, h=10, text=f"Por completar el curso de {self.curso}", align="C")
        pdf.ln(25)

        pdf.set_font("Arial", style="I", size=14)
        pdf.cell(w=0, h=10, text=f"Fecha: {self.fecha}", align="C")

        return pdf

    def nombre_archivo_seguro(self, nombre_alumno):
        """Convierte el nombre en un nombre de archivo seguro."""
        seguro = (
            nombre_alumno
            .replace(" ", "_")
            .replace("á", "a").replace("é", "e").replace("í", "i")
            .replace("ó", "o").replace("ú", "u").replace("ñ", "n")
            .replace("Á", "A").replace("É", "E").replace("Í", "I")
            .replace("Ó", "O").replace("Ú", "U").replace("Ñ", "N")
        )
        return f"certificado_{seguro}.pdf"

    def guardar(self, pdf, nombre_archivo):
        ruta = self.carpeta_salida / nombre_archivo
        pdf.output(str(ruta))
        return ruta

    def generar_desde_excel(self, nombre_excel="alumnos.xlsx", fecha=""):
        """Lee el Excel y crea un PDF por cada alumno."""
        if not fecha:
            fecha = input("Fecha para todos los certificados: ")

        alumnos = self.leer_excel(nombre_excel)
        print(f"\n=== Leidos {len(alumnos)} alumnos de {nombre_excel} ===\n")

        for alumno in alumnos:
            self.nombre = alumno["nombre"]
            self.curso = alumno["curso"]
            self.fecha = fecha

            pdf = self.crear_pdf()
            archivo = self.nombre_archivo_seguro(self.nombre)
            ruta = self.guardar(pdf, archivo)
            print(f"OK → {ruta.name}  ({self.nombre})")

        print(f"\nListo! Revisa la carpeta:\n{self.carpeta_salida}")


if __name__ == "__main__":
    generador = GeneradorCertificado()
    generador.generar_desde_excel("alumnos.xlsx")
