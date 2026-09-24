# ============================================================
# Aplicación con página web: generador de certificados
# ============================================================
# La persona no escribe código. Doble clic abre el navegador
# con una página hecha en HTML y una hoja de estilos (CSS).
# El dibujo del PDF sigue siendo el del paso 9.
#
# Cómo abrirla:
#   Doble clic en "Abrir generador.command"
#   o, desde esta carpeta:
#   .venv/bin/python app_certificados.py
# ============================================================

import json
import socket
import subprocess
import sys
import threading
import traceback
import webbrowser
from datetime import date
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

from openpyxl import load_workbook

from paso9_excel_con_borde import GeneradorCertificado


CARPETA = Path(__file__).resolve().parent
INTERFAZ = CARPETA / "interfaz"
PAGINAS = {
    "/": ("index.html", "text/html; charset=utf-8"),
    "/estilos.css": ("estilos.css", "text/css; charset=utf-8"),
    "/app.js": ("app.js", "text/javascript; charset=utf-8"),
}

MESES = (
    "",
    "enero",
    "febrero",
    "marzo",
    "abril",
    "mayo",
    "junio",
    "julio",
    "agosto",
    "septiembre",
    "octubre",
    "noviembre",
    "diciembre",
)


def fecha_de_hoy():
    """Devuelve la fecha de hoy en español, lista para el certificado."""
    hoy = date.today()
    return f"{hoy.day} de {MESES[hoy.month]} de {hoy.year}"


class GeneradorDesdeArchivo(GeneradorCertificado):
    """Misma fábrica del paso 9, pero con archivos que elige la persona."""

    def __init__(self):
        self.carpeta = CARPETA
        self.fuente_normal = "/System/Library/Fonts/Supplemental/Arial.ttf"
        self.fuente_negrita = "/System/Library/Fonts/Supplemental/Arial Bold.ttf"
        self.fuente_cursiva = "/System/Library/Fonts/Supplemental/Arial Italic.ttf"
        self.nombre = ""
        self.curso = ""
        self.fecha = ""
        self.carpeta_salida = self.carpeta

    def leer_desde(self, ruta):
        """
        Lee nombre y curso desde cualquier Excel.
        Devuelve (alumnos, avisos). Los avisos son filas que se omitieron.
        """
        ruta = Path(ruta)
        if not ruta.exists():
            raise FileNotFoundError(f"No encuentro el archivo:\n{ruta}")

        try:
            libro = load_workbook(ruta, data_only=True)
        except Exception as error:
            raise ValueError(
                "No pude abrir ese archivo. Tiene que ser un Excel .xlsx "
                "con las columnas nombre y curso."
            ) from error

        hoja = libro.active
        encabezados = [(c.value or "").strip().lower() for c in hoja[1]]
        try:
            col_nombre = encabezados.index("nombre") + 1
            col_curso = encabezados.index("curso") + 1
        except ValueError as error:
            raise ValueError(
                "La primera fila del Excel debe tener las columnas: nombre y curso."
            ) from error

        alumnos = []
        avisos = []
        for fila in range(2, hoja.max_row + 1):
            nombre = hoja.cell(row=fila, column=col_nombre).value
            curso = hoja.cell(row=fila, column=col_curso).value
            nombre = str(nombre).strip() if nombre else ""
            curso = str(curso).strip() if curso else ""
            if not nombre and not curso:
                continue
            if not nombre or not curso:
                avisos.append(f"Fila {fila}: falta el nombre o el curso. Se omitió.")
                continue
            alumnos.append({"nombre": nombre, "curso": curso})

        if not alumnos:
            raise ValueError(
                "El Excel no tiene alumnos. Agrega filas debajo de nombre y curso."
            )
        return alumnos, avisos

    def generar_archivos(self, ruta_excel, fecha, carpeta_salida):
        """Crea un PDF por alumno y devuelve las rutas y los avisos."""
        for fuente in (self.fuente_normal, self.fuente_negrita, self.fuente_cursiva):
            if not Path(fuente).exists():
                raise FileNotFoundError(
                    "No encuentro la fuente Arial de este Mac, "
                    "y el certificado la necesita para escribir los acentos."
                )

        carpeta_salida = Path(carpeta_salida)
        carpeta_salida.mkdir(parents=True, exist_ok=True)
        self.carpeta_salida = carpeta_salida

        alumnos, avisos = self.leer_desde(ruta_excel)
        self.fecha = fecha.strip()
        rutas = []
        usados = set()

        for alumno in alumnos:
            self.nombre = alumno["nombre"]
            self.curso = alumno["curso"]
            pdf = self.crear_pdf()
            archivo = self._nombre_sin_repetir(self.nombre, usados)
            rutas.append(self.guardar(pdf, archivo))

        return rutas, avisos

    def _nombre_sin_repetir(self, nombre_alumno, usados):
        """Si dos alumnos se llaman igual, el segundo PDF no pisa al primero."""
        archivo = self.nombre_archivo_seguro(nombre_alumno)
        if archivo not in usados:
            usados.add(archivo)
            return archivo

        base = archivo[:-4]
        numero = 2
        while f"{base}_{numero}.pdf" in usados:
            numero += 1
        archivo = f"{base}_{numero}.pdf"
        usados.add(archivo)
        return archivo


class Estado:
    """Recuerda la última carpeta para poder abrirla en el Finder."""

    def __init__(self):
        self.generador = GeneradorDesdeArchivo()
        self.candado = threading.Lock()
        self.carpeta = None


ESTADO = Estado()


def dialogo_mac(orden):
    """Abre el selector de archivos del Mac y devuelve la ruta, o None si cancela."""
    resultado = subprocess.run(
        ["osascript", "-e", orden],
        capture_output=True,
        text=True,
    )
    if resultado.returncode != 0:
        return None
    return resultado.stdout.strip() or None


def abrir_en_el_sistema(carpeta):
    if sys.platform == "darwin":
        subprocess.run(["open", str(carpeta)], check=False)
    elif sys.platform == "win32":
        subprocess.run(["explorer", str(carpeta)], check=False)
    else:
        subprocess.run(["xdg-open", str(carpeta)], check=False)


class Manejador(BaseHTTPRequestHandler):
    """Sirve la página y responde a los botones."""

    def log_message(self, formato, *args):
        return

    def do_GET(self):
        if self.path == "/api/inicio":
            self._enviar_json(200, {"fecha": fecha_de_hoy()})
            return

        pagina = PAGINAS.get(self.path)
        if pagina is None:
            self._enviar_texto(404, "No encontrado")
            return

        nombre, tipo = pagina
        archivo = INTERFAZ / nombre
        if not archivo.is_file():
            self._enviar_texto(404, "No encontrado")
            return

        contenido = archivo.read_bytes()
        self.send_response(200)
        self.send_header("Content-Type", tipo)
        self.send_header("Content-Length", str(len(contenido)))
        self.send_header("Cache-Control", "no-cache")
        self.end_headers()
        self.wfile.write(contenido)

    def do_POST(self):
        if self.path == "/api/elegir-excel":
            self._elegir_excel()
        elif self.path == "/api/elegir-carpeta":
            self._elegir_carpeta()
        elif self.path == "/api/generar":
            self._generar()
        elif self.path == "/api/abrir-carpeta":
            self._abrir_carpeta()
        else:
            self._enviar_texto(404, "No encontrado")

    def _elegir_excel(self):
        ruta = dialogo_mac(
            'POSIX path of (choose file with prompt "Elige el Excel de alumnos")'
        )
        if not ruta:
            self._enviar_json(200, {"cancelado": True})
            return
        if Path(ruta).suffix.lower() != ".xlsx":
            self._enviar_json(400, {
                "ok": False,
                "error": "Elige un archivo Excel .xlsx.",
            })
            return
        sugerida = str(Path(ruta).parent / "certificados")
        self._enviar_json(200, {"ruta": ruta, "carpetaSugerida": sugerida})

    def _elegir_carpeta(self):
        ruta = dialogo_mac(
            'POSIX path of (choose folder with prompt "Elige dónde guardar los PDF")'
        )
        if not ruta:
            self._enviar_json(200, {"cancelado": True})
            return
        ruta = str(Path(ruta))
        ESTADO.carpeta = Path(ruta)
        self._enviar_json(200, {"ruta": ruta})

    def _generar(self):
        try:
            datos = self._leer_json()
        except ValueError as error:
            self._enviar_json(400, {"ok": False, "error": str(error)})
            return

        excel = str(datos.get("excel", "")).strip()
        fecha = str(datos.get("fecha", "")).strip()
        carpeta = str(datos.get("carpeta", "")).strip()
        if not excel:
            self._enviar_json(400, {
                "ok": False,
                "error": "Elige el archivo Excel con los alumnos.",
            })
            return
        if not fecha:
            self._enviar_json(400, {
                "ok": False,
                "error": "Escribe la fecha que debe aparecer en los certificados.",
            })
            return
        if not carpeta:
            self._enviar_json(400, {
                "ok": False,
                "error": "Elige la carpeta donde se guardarán los PDF.",
            })
            return

        try:
            with ESTADO.candado:
                rutas, avisos = ESTADO.generador.generar_archivos(excel, fecha, carpeta)
        except (ValueError, FileNotFoundError) as error:
            self._enviar_json(400, {"ok": False, "error": str(error)})
            return
        except Exception:
            traceback.print_exc()
            self._enviar_json(500, {
                "ok": False,
                "error": "No se pudieron crear los certificados.",
            })
            return

        ESTADO.carpeta = Path(carpeta)
        self._enviar_json(200, {
            "ok": True,
            "archivos": [ruta.name for ruta in rutas],
            "avisos": avisos,
            "carpeta": carpeta,
        })

    def _abrir_carpeta(self):
        if ESTADO.carpeta is None or not ESTADO.carpeta.is_dir():
            self._enviar_json(400, {
                "ok": False,
                "error": "Todavía no hay una carpeta de certificados para abrir.",
            })
            return
        abrir_en_el_sistema(ESTADO.carpeta)
        self._enviar_json(200, {"ok": True})

    def _leer_json(self):
        longitud = int(self.headers.get("Content-Length", "0") or "0")
        if longitud > 1_000_000:
            raise ValueError("La petición es demasiado grande.")
        if longitud == 0:
            return {}
        cuerpo = self.rfile.read(longitud)
        try:
            datos = json.loads(cuerpo.decode("utf-8"))
        except json.JSONDecodeError as error:
            raise ValueError("No pude leer los datos enviados.") from error
        if not isinstance(datos, dict):
            raise ValueError("No pude leer los datos enviados.")
        return datos

    def _enviar_json(self, codigo, datos):
        cuerpo = json.dumps(datos, ensure_ascii=False).encode("utf-8")
        self.send_response(codigo)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(cuerpo)))
        self.end_headers()
        self.wfile.write(cuerpo)

    def _enviar_texto(self, codigo, texto):
        cuerpo = texto.encode("utf-8")
        self.send_response(codigo)
        self.send_header("Content-Type", "text/plain; charset=utf-8")
        self.send_header("Content-Length", str(len(cuerpo)))
        self.end_headers()
        self.wfile.write(cuerpo)


def puerto_libre():
    with socket.socket() as conexion:
        conexion.bind(("127.0.0.1", 0))
        return conexion.getsockname()[1]


def main():
    puerto = puerto_libre()
    servidor = ThreadingHTTPServer(("127.0.0.1", puerto), Manejador)
    direccion = f"http://127.0.0.1:{puerto}"
    print("Generador de certificados")
    print(f"La página se abrió en el navegador.\nSi no se abrió, entra aquí: {direccion}")
    print("\nDeja esta ventana abierta mientras uses el programa.")
    print("Para terminarlo, pulsa Control+C.")
    webbrowser.open(direccion)
    try:
        servidor.serve_forever()
    except KeyboardInterrupt:
        print("\nPrograma cerrado.")
    finally:
        servidor.server_close()


if __name__ == "__main__":
    main()
