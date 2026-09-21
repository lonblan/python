# ============================================================
# PASO 9 (FINAL): Excel + borde decorativo
# ============================================================
# Une lo mejor de los pasos anteriores:
#   - Lee alumnos desde alumnos.xlsx  (paso 7)
#   - Dibuja doble borde en cada PDF   (paso 8)
#   - Un archivo por alumno            (pasos 6/7)
#
# Cómo usarlo:
#   1) Edita alumnos.xlsx (columnas: nombre | curso)
#   2) Ejecuta este script
#   3) Revisa la carpeta certificados_final/
# ============================================================

from pathlib import Path
from fpdf import FPDF
from openpyxl import load_workbook


class GeneradorCertificado:
    """Certificados desde Excel, con borde decorativo."""

    def __init__(self):
        self.carpeta = Path(__file__).resolve().parent
        self.carpeta_salida = self.carpeta / "certificados_final"
        self.carpeta_salida.mkdir(exist_ok=True)

        self.fuente_normal = "/System/Library/Fonts/Supplemental/Arial.ttf"
        self.fuente_negrita = "/System/Library/Fonts/Supplemental/Arial Bold.ttf"
        self.fuente_cursiva = "/System/Library/Fonts/Supplemental/Arial Italic.ttf"

        self.nombre = ""
        self.curso = ""
        self.fecha = ""

    def leer_excel(self, nombre_archivo="alumnos.xlsx"):
        """Lee nombre y curso desde el Excel."""
        ruta = self.carpeta / nombre_archivo

        if not ruta.exists():
            raise FileNotFoundError(
                f"No encuentro el archivo:\n{ruta}\n"
                "Pon alumnos.xlsx en la misma carpeta que este script."
            )

        libro = load_workbook(ruta)
        hoja = libro.active

        encabezados = [(c.value or "").strip().lower() for c in hoja[1]]
        try:
            col_nombre = encabezados.index("nombre") + 1
            col_curso = encabezados.index("curso") + 1
        except ValueError:
            raise ValueError(
                'La primera fila debe tener las columnas: nombre | curso'
            )

        alumnos = []
        for fila in range(2, hoja.max_row + 1):
            nombre = hoja.cell(row=fila, column=col_nombre).value
            curso = hoja.cell(row=fila, column=col_curso).value
            if not nombre and not curso:
                continue
            alumnos.append({
                "nombre": str(nombre).strip(),
                "curso": str(curso).strip() if curso else "",
            })
        return alumnos

    def dibujar_borde(self, pdf):
        """Doble marco alrededor de la página."""
        ancho = pdf.w
        alto = pdf.h

        pdf.set_draw_color(25, 55, 110)
        pdf.set_line_width(1.8)
        pdf.rect(x=10, y=10, w=ancho - 20, h=alto - 20)

        pdf.set_line_width(0.6)
        pdf.rect(x=14, y=14, w=ancho - 28, h=alto - 28)

    def crear_pdf(self):
        """Crea un certificado con borde y texto."""
        pdf = FPDF(orientation="L", unit="mm", format="A4")
        pdf.add_page()

        self.dibujar_borde(pdf)

        pdf.add_font("Arial", style="", fname=self.fuente_normal)
        pdf.add_font("Arial", style="B", fname=self.fuente_negrita)
        pdf.add_font("Arial", style="I", fname=self.fuente_cursiva)

        pdf.set_y(40)

        pdf.set_font("Arial", style="B", size=28)
        pdf.cell(w=0, h=20, text="CERTIFICADO DE CURSO", align="C")
        pdf.ln(28)

        pdf.set_font("Arial", style="", size=16)
        pdf.cell(w=0, h=10, text="Se otorga el presente certificado a:", align="C")
        pdf.ln(20)

        pdf.set_font("Arial", style="B", size=24)
        pdf.cell(w=0, h=12, text=self.nombre, align="C")
        pdf.ln(20)

        pdf.set_font("Arial", style="", size=16)
        pdf.cell(w=0, h=10, text=f"Por completar el curso de {self.curso}", align="C")
        pdf.ln(25)

        pdf.set_font("Arial", style="I", size=14)
        pdf.cell(w=0, h=10, text=f"Fecha: {self.fecha}", align="C")

        return pdf

    def nombre_archivo_seguro(self, nombre_alumno):
        seguro = (
            nombre_alumno
            .replace(" ", "_")
            .replace("á", "a").replace("é", "e").replace("í", "i")
            .replace("ó", "o").replace("ú", "u").replace("ñ", "n")
            .replace("Á", "A").replace("É", "E").replace("Í", "I")
            .replace("Ó", "O").replace("Ú", "U").replace("Ñ", "N")
        )
        return f"certificado_{seguro}.pdf"

    def guardar(self, pdf, nombre_archivo):
        ruta = self.carpeta_salida / nombre_archivo
        pdf.output(str(ruta))
        return ruta

    def generar_desde_excel(self, nombre_excel="alumnos.xlsx", fecha=""):
        """Lee el Excel y crea un PDF con borde por cada alumno."""
        if not fecha:
            fecha = input("Fecha para todos los certificados: ")

        alumnos = self.leer_excel(nombre_excel)
        print(f"\n=== Generando {len(alumnos)} certificados (Excel + borde) ===\n")

        for alumno in alumnos:
            self.nombre = alumno["nombre"]
            self.curso = alumno["curso"]
            self.fecha = fecha

            pdf = self.crear_pdf()
            archivo = self.nombre_archivo_seguro(self.nombre)
            ruta = self.guardar(pdf, archivo)
            print(f"OK → {ruta.name}  ({self.nombre})")

        print(f"\nListo! Revisa la carpeta:\n{self.carpeta_salida}")


if __name__ == "__main__":
    generador = GeneradorCertificado()
    generador.generar_desde_excel("alumnos.xlsx")
